#!/usr/bin/env python
"""硬件事实报告 Excel 生成（格式按 assets/circuit_facts_example.xlsx，
即《GD32F205VET_pin_config.xlsx》风格）。

Sheet 1 "<MCU> 引脚配置"：
- 列：引脚号 | 引脚名 | 网络 | 作用 | 外设/模式 | 状态
- 表头蓝底白字；数据行按「外设/模式」类别分色（与示例图例一致）：
  浅蓝=电源/地/参考  黄=复位/晶振/BOOT/调试  绿=数字输入
  橙=模拟输入 ADC    浅蓝=GPIO 输出/UART       紫=SPI/I2C
  粉红=定时器 PWM    灰=FSMC/EXMC             无填充=NC
- 状态列：ok 留空；warning=待确认；conflict=冲突；nc=NC

Sheet 2 "图例"：颜色分组说明 + 状态说明（与示例结构一致）。
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from pin_inferrer import _COLOR_BY_PERIPHERAL

HEADER_BG = "4472C4"   # 表头深蓝（与示例一致）
HEADER_FONT = Font(color="FFFFFF", bold=True)

# 类别 → 填充色（与 GD32 示例图例一致）
CATEGORY_COLORS = {
    "power": "D9E1F2",       # 电源/地/参考
    "special": "FFF2CC",     # 复位/晶振/BOOT/调试
    "digital_in": "E2EFDA",  # 数字输入
    "analog": "FCE4D6",      # 模拟输入 ADC
    "output_uart": "DDEBF7", # GPIO 输出 / UART
    "spi_i2c": "E4DFEC",     # SPI / I2C
    "pwm": "F2DCDB",         # 定时器 PWM
    "exmc": "EDEDED",        # FSMC / EXMC
}

STATUS_TEXT = {"ok": "", "warning": "待确认", "conflict": "冲突", "nc": "NC"}
STATUS_FONT = {
    "conflict": Font(color="FF0000", bold=True),
    "warning": Font(color="BF8F00", bold=True),
    "nc": Font(color="808080"),
}

_COLUMNS = ["引脚号", "引脚名", "网络", "作用", "外设/模式", "状态"]
_WIDTHS = [8, 18, 16, 26, 16, 8]

_LEGEND_COLOR_ROWS = [
    ("浅蓝", "power", "电源/地/参考"),
    ("黄色", "special", "复位/晶振/BOOT/调试"),
    ("绿色", "digital_in", "数字输入（按键/旋钮/开关）"),
    ("橙色", "analog", "模拟输入 ADC"),
    ("浅蓝", "output_uart", "GPIO 输出控制 / UART 串口"),
    ("紫色", "spi_i2c", "SPI / I2C"),
    ("粉红", "pwm", "定时器 PWM"),
    ("灰色", "exmc", "FSMC / EXMC 并口"),
    ("无填充", None, "未连接（NC）"),
]


def _category_of(peripheral: str) -> str | None:
    """外设/模式 → 颜色类别。NC/空 → None（无填充）。"""
    if not peripheral or peripheral == "NC":
        return None
    for pat, cat in _COLOR_BY_PERIPHERAL:
        if pat.search(peripheral):
            return cat
    return "digital_in"  # 兜底：按数字信号着色


def export_facts_xlsx(facts: dict, out_path: Path) -> None:
    """生成 circuit_facts.xlsx。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{facts.get('mcu') or facts.get('platform') or 'MCU'} 引脚配置"

    # ---- 表头 ----
    for col, title in enumerate(_COLUMNS, 1):
        c = ws.cell(row=1, column=col, value=title)
        c.fill = PatternFill("solid", fgColor=HEADER_BG)
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center")
    for col, w in enumerate(_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A2"

    # ---- 数据行 ----
    for i, p in enumerate(facts.get("pins", []), start=2):
        row_vals = [
            int(p["pin"]) if str(p["pin"]).isdigit() else p["pin"],
            p["pin_name"],
            p.get("net") or "-",
            p.get("role") or "-",
            p.get("peripheral") or "-",
            STATUS_TEXT.get(p.get("status", "ok"), ""),
        ]
        cat = _category_of(p.get("peripheral") or "")
        for col, val in enumerate(row_vals, 1):
            c = ws.cell(row=i, column=col, value=val)
            if cat:
                c.fill = PatternFill("solid", fgColor=CATEGORY_COLORS[cat])
            if col == len(row_vals) and p.get("status") in STATUS_FONT:
                c.font = STATUS_FONT[p["status"]]

    # ---- 图例 sheet ----
    lg = wb.create_sheet("图例")
    lg.cell(row=1, column=1, value="颜色分组").font = Font(bold=True)
    for r, (color_name, cat, desc) in enumerate(_LEGEND_COLOR_ROWS, start=2):
        c1 = lg.cell(row=r, column=1, value=color_name)
        if cat:
            c1.fill = PatternFill("solid", fgColor=CATEGORY_COLORS[cat])
        lg.cell(row=r, column=2, value=desc)
    base = len(_LEGEND_COLOR_ROWS) + 3
    lg.cell(row=base, column=1, value="状态说明").font = Font(bold=True)
    for r, (status, desc) in enumerate(
        [
            ("待确认", "网络名无明确语义，无法推断功能角色，建议人工/AI 补充"),
            ("冲突", "存在复用冲突或电源缺失等 error 级问题，必须修复"),
            ("NC", "网表中无连接的引脚"),
        ],
        start=base + 1,
    ):
        lg.cell(row=r, column=1, value=status).font = STATUS_FONT.get(
            "warning" if status == "待确认" else "conflict" if status == "冲突" else "nc"
        )
        lg.cell(row=r, column=2, value=desc)
    lg.column_dimensions["A"].width = 10
    lg.column_dimensions["B"].width = 52

    wb.save(out_path)
