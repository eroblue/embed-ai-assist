#!/usr/bin/env python
"""从通用网表数据生成引脚配置 Excel（outputs/pin_table.xlsx）。

格式规范见 references/excel_format.md：
- 所有元件放在同一个工作表中，每个元件一个区块，便于通读
- 区块 = 元件标题行（合并单元格）+ 表头行 + 引脚行 + 空行分隔
- 列：引脚号 / 引脚名 / 网络名（作用与外设/模式由 S4 circuit-investigator 判断，
  输出在 circuit_facts.xlsx，本表不再输出这两列）
- 行颜色：电源=蓝、特殊功能=黄、悬空=灰、普通=白
"""

from __future__ import annotations

import re
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:  # 由调用方处理
    Workbook = None  # type: ignore[assignment]

SHEET_TITLE = "引脚配置"

HEADER = ["引脚号", "引脚名", "网络名"]

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
COMP_FILL = PatternFill(start_color="B4C6E7", end_color="B4C6E7", fill_type="solid")
COMP_FONT = Font(bold=True)
FILL_POWER = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
FILL_SPECIAL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
FILL_FLOATING = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
COL_WIDTHS = [10, 14, 18]

# 数字电源网络名：3V3 / 5V / 12V / 2V8 / +3V3 / 3.3V 等（整体匹配）
_POWER_NET_RE = re.compile(r"^\+?\d+(\.\d+)?V\d*$")


def infer_role(pin_name: str, net: str | None, pin_type: str) -> str:
    """按引脚名/网络名/引脚类型推断行类别（仅用于行颜色），无法识别返回空串。

    语义来源说明：原理图只存电气连接，"按键/LED/串口"等用途语义藏在网络名中
    （如 KEY2 / LED1 / UART1_TX），引脚级复用映射则需 datasheet（S3/S4 的职责）。
    """
    name_l = (pin_name or "").upper()
    net_l = (net or "").upper()

    if pin_type == "POWER" or any(
        k in name_l or k in net_l
        for k in ("VDD", "VSS", "VCC", "GND", "VBAT", "VREF")
    ):
        return "电源"
    if net and _POWER_NET_RE.match(net_l):
        return "电源"
    if any(k in name_l or k in net_l for k in ("SWDIO", "SWCLK", "TMS", "TCK")):
        return "SWD 调试"
    if any(k in net_l for k in ("JTDI", "JTDO", "JTRST")):
        return "JTAG 调试"
    if "BOOT" in name_l or "BOOT" in net_l:
        return "启动模式配置"
    if any(k in name_l or k in net_l for k in ("NRST", "RESET", "RST")):
        return "复位"
    if any(k in name_l for k in ("OSCI", "OSCO", "XCIN", "XCOUT", "XTAL")) or "OSC" in net_l:
        return "晶振"
    if any(k in net_l for k in ("KEY", "BTN", "BUTTON")):
        return "按键输入"
    if "LED" in net_l:
        return "LED 指示"
    if net_l in ("SCL", "SDA") or "I2C" in net_l:
        return "I2C 通信"
    if any(k in net_l for k in ("MOSI", "MISO", "SPI")):
        return "SPI 通信"
    if "CAN" in net_l:
        return "CAN 通信"
    if "UART" in net_l or "USART" in net_l:
        return "串口通信"
    if net is None:
        return "未使用（悬空）"
    return ""


def _row_fill(role: str, net: str | None) -> PatternFill | None:
    """根据推断结果选择行填充色。"""
    if role == "电源":
        return FILL_POWER
    if role in (
        "SWD 调试", "JTAG 调试", "启动模式配置", "复位", "晶振",
        "I2C 通信", "SPI 通信", "CAN 通信", "串口通信",
    ):
        return FILL_SPECIAL
    if net is None:
        return FILL_FLOATING
    return None


def export_pin_table(netlist_doc: dict, excel_path: Path) -> tuple[int, list[str]]:
    """生成引脚配置 Excel，返回 (引脚总行数, 警告列表)。

    netlist_doc: parse.py 生成的网表文档（含 components）
    excel_path: 输出文件路径
    """
    if Workbook is None:
        raise ImportError("openpyxl 未安装: pip install openpyxl")

    warnings: list[str] = []
    components = netlist_doc.get("components", [])
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE
    for col_idx, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[chr(64 + col_idx)].width = width

    total_rows = 0
    row = 1
    for comp in components:
        designator = comp.get("designator") or "?"
        value = (comp.get("value") or "").strip()
        pins = comp.get("pins", [])

        # 元件标题行（合并 A:C）
        title = f"{designator}（{value}）— {len(pins)} 引脚" if value else f"{designator} — {len(pins)} 引脚"
        ws.cell(row=row, column=1, value=title)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        for col in range(1, 4):
            cell = ws.cell(row=row, column=col)
            cell.fill = COMP_FILL
            cell.font = COMP_FONT
            cell.alignment = ALIGN_LEFT
        row += 1

        # 表头行
        for col, head in enumerate(HEADER, start=1):
            cell = ws.cell(row=row, column=col, value=head)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = ALIGN_CENTER
        row += 1

        # 引脚数据行
        for pin in pins:
            net = pin.get("net")
            role = infer_role(
                pin.get("pin_name", ""), net, pin.get("pin_type", "")
            )
            ws.cell(row=row, column=1, value=pin.get("pin", ""))
            ws.cell(row=row, column=2, value=pin.get("pin_name", ""))
            ws.cell(row=row, column=3, value=net if net is not None else "未连接")
            fill = _row_fill(role, net)
            if fill is not None:
                for col in range(1, 4):
                    ws.cell(row=row, column=col).fill = fill
            ws.cell(row=row, column=1).alignment = ALIGN_CENTER
            row += 1
            total_rows += 1

        row += 1  # 元件区块间空行

    if total_rows == 0:
        warnings.append("网表无元件引脚，pin_table.xlsx 仅含空表")

    excel_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(excel_path)
    return total_rows, warnings
