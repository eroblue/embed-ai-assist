#!/usr/bin/env python
"""从芯片提取数据生成两类 Excel（outputs/chip_info/）。

- pin_table.xlsx：引脚定义表（单工作表平铺，144 行量级便于通读筛选）
- register_map.xlsx：寄存器映射表（每个外设一个工作表，外设寄存器多，
  按外设分组查阅）

格式规范见 references/excel_format.md。
"""

from __future__ import annotations

import re
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    Workbook = None  # type: ignore[assignment]

PIN_HEADER = ["引脚号", "引脚名", "类型", "复用功能（默认）", "复用功能（重映射）", "电气特性"]
REG_HEADER = ["外设", "寄存器名", "地址", "偏移", "位域", "读写权限", "复位值", "功能说明"]

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
FILL_POWER = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
FILL_SPECIAL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
PIN_COL_WIDTHS = [10, 12, 10, 40, 30, 10]
REG_COL_WIDTHS = [14, 22, 14, 10, 46, 10, 14, 40]


def _pin_type_label(p: dict) -> str:
    return {
        "IO": "IO",
        "POWER": "电源",
        "INPUT": "输入",
        "OUTPUT": "输出",
        "BOOT": "启动",
        "NRST": "复位",
        "UNKNOWN": "未知",
    }.get(p.get("pin_type", "UNKNOWN"), "未知")


def export_pin_table(pins_doc: dict, excel_path: Path) -> tuple[int, list[str]]:
    """生成引脚定义 Excel（单工作表平铺）。"""
    if Workbook is None:
        raise ImportError("openpyxl 未安装: pip install openpyxl")

    warnings: list[str] = []
    wb = Workbook()
    ws = wb.active
    ws.title = "引脚定义"
    for col, width in enumerate(PIN_COL_WIDTHS, start=1):
        ws.column_dimensions[chr(64 + col)].width = width

    # 标题行
    ws.cell(row=1, column=1, value=f"{pins_doc.get('package', '')} 引脚定义（共 {pins_doc.get('pin_count', 0)} 脚）")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(PIN_HEADER))
    c = ws.cell(row=1, column=1)
    c.font = Font(bold=True)
    c.alignment = ALIGN_LEFT

    for col, head in enumerate(PIN_HEADER, start=1):
        cell = ws.cell(row=2, column=col, value=head)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER

    row = 3
    for p in pins_doc.get("pins", []):
        has_af = bool(p.get("alternate_functions")) or bool(p.get("alternate_functions_remap"))
        ws.cell(row=row, column=1, value=p.get("pin"))
        ws.cell(row=row, column=2, value=p.get("name"))
        ws.cell(row=row, column=3, value=_pin_type_label(p))
        ws.cell(row=row, column=4, value=" / ".join(p.get("alternate_functions", [])))
        ws.cell(row=row, column=5, value=" / ".join(p.get("alternate_functions_remap", [])))
        ws.cell(row=row, column=6, value=p.get("io_level", ""))
        # 颜色：电源=蓝底，含复用=黄底，普通=白底
        if p.get("pin_type") == "POWER":
            fill = FILL_POWER
        elif has_af:
            fill = FILL_SPECIAL
        else:
            fill = None
        if fill is not None:
            for col in range(1, len(PIN_HEADER) + 1):
                ws.cell(row=row, column=col).fill = fill
        ws.cell(row=row, column=1).alignment = ALIGN_CENTER
        row += 1

    if row == 3:
        warnings.append("pins 数据为空，pin_table.xlsx 仅含表头")

    excel_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(excel_path)
    return row - 3, warnings


def export_register_map(registers_doc: dict, excel_path: Path) -> tuple[int, list[str]]:
    """生成寄存器映射 Excel（每个外设一个工作表）。"""
    if Workbook is None:
        raise ImportError("openpyxl 未安装: pip install openpyxl")

    warnings: list[str] = []
    registers = registers_doc.get("registers", [])
    by_periph: dict[str, list] = {}
    for r in registers:
        by_periph.setdefault(r.get("peripheral", "?"), []).append(r)

    wb = Workbook()
    wb.remove(wb.active)
    rows_written = 0
    used_titles: set[str] = set()
    for periph in sorted(by_periph):
        # Excel 工作表名上限 31 字符，且不允许 / \ ? * [ ] :
        title = re.sub(r"[/\\?*\[\]:]", "_", periph).strip()[:31] or "未分组"
        if title in used_titles:  # 截断/替换后可能重名，追加序号
            n = 2
            while f"{title[:28]}_{n}" in used_titles:
                n += 1
            title = f"{title[:28]}_{n}"
        used_titles.add(title)
        ws = wb.create_sheet(title=title)
        for col, width in enumerate(REG_COL_WIDTHS, start=1):
            ws.column_dimensions[chr(64 + col)].width = width
        for col, head in enumerate(REG_HEADER, start=1):
            cell = ws.cell(row=1, column=col, value=head)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = ALIGN_CENTER
        row = 2
        for r in by_periph[periph]:
            bit_summary = "; ".join(
                f"{bf.get('name')}[{bf.get('bits')}]" for bf in r.get("bitfields", [])
            )
            ws.cell(row=row, column=1, value=r.get("peripheral"))
            ws.cell(row=row, column=2, value=r.get("name"))
            ws.cell(row=row, column=3, value=r.get("address"))
            ws.cell(row=row, column=4, value=r.get("offset"))
            ws.cell(row=row, column=5, value=bit_summary)
            ws.cell(row=row, column=6, value=r.get("access", ""))
            ws.cell(row=row, column=7, value=r.get("reset_value") or "")
            ws.cell(row=row, column=8, value=r.get("description", ""))
            row += 1
            rows_written += 1

    if rows_written == 0:
        warnings.append("registers 数据为空，register_map.xlsx 未生成内容")

    excel_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(excel_path)
    return rows_written, warnings
